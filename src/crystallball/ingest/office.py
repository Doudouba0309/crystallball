"""Loaders for Office Open XML formats: .docx, .pptx, and .xlsx."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterator

from ..errors import LoadError
from ..models import (
    BLOCK_HEADING,
    BLOCK_PARAGRAPH,
    BLOCK_SHEET_ROW,
    BLOCK_SLIDE,
    BLOCK_TABLE_ROW,
    Document,
    SourceRef,
    TextBlock,
)
from ._optional import require

#: Matches built-in Word heading styles such as "Heading 1".
_HEADING_RE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


def _docx_heading_level(paragraph: Any) -> int | None:
    """Return the heading level of a Word paragraph, if it is a heading."""
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", "") or ""
    match = _HEADING_RE.search(name)
    if match:
        return int(match.group(1))
    if name.strip().lower() == "title":
        return 1
    return None


class DocxLoader:
    """Loader for Word documents, preserving paragraph and table order."""

    name = "docx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    def load(self, path: Path, source: SourceRef) -> Document:
        docx = require("docx", "office", purpose="reading .docx files")
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # noqa: BLE001 - library raises varied types
            raise LoadError(f"could not open Word document {path}: {exc}") from exc

        blocks: list[TextBlock] = []
        table_index = 0
        paragraph_index = 0

        for item in _iter_docx_blocks(document, docx):
            if hasattr(item, "rows"):  # a Table
                table_index += 1
                for row_number, row in enumerate(item.rows, start=1):
                    cells = [
                        " ".join(cell.text.split()) for cell in row.cells
                    ]
                    blocks.append(
                        TextBlock(
                            kind=BLOCK_TABLE_ROW,
                            text=" | ".join(c for c in cells if c),
                            locator=f"table{table_index}!row{row_number}",
                        )
                    )
                continue

            text = " ".join(item.text.split())
            if not text:
                continue
            level = _docx_heading_level(item)
            if level is not None:
                blocks.append(
                    TextBlock(
                        kind=BLOCK_HEADING, text=text, locator="heading", level=level
                    )
                )
            else:
                paragraph_index += 1
                blocks.append(
                    TextBlock(
                        kind=BLOCK_PARAGRAPH, text=text, locator=f"p.{paragraph_index}"
                    )
                )

        return Document(
            source=source,
            text="\n\n".join(b.text for b in blocks),
            blocks=blocks,
            metadata={
                "table_count": table_index,
                "paragraph_count": paragraph_index,
                "heading_count": sum(1 for b in blocks if b.kind == BLOCK_HEADING),
                **_docx_properties(document),
            },
        )


def _iter_docx_blocks(document: Any, docx_module: Any) -> Iterator[Any]:
    """Yield paragraphs and tables in document order.

    Prefers python-docx's own iterator, falling back to walking the XML body
    directly for older versions.
    """
    if hasattr(document, "iter_inner_content"):
        yield from document.iter_inner_content()
        return

    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _docx_properties(document: Any) -> dict[str, Any]:
    """Core document properties, normalised to primitives."""
    props = getattr(document, "core_properties", None)
    if props is None:
        return {}
    result: dict[str, Any] = {}
    for key in ("title", "author", "subject", "category", "comments"):
        value = getattr(props, key, None)
        if value:
            result[f"docx_{key}"] = str(value)
    for key in ("created", "modified"):
        value = getattr(props, key, None)
        if value is not None:
            result[f"docx_{key}"] = value.isoformat()
    return result


class PptxLoader:
    """Loader for PowerPoint decks, one block per text frame or table row."""

    name = "pptx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    def load(self, path: Path, source: SourceRef) -> Document:
        pptx = require("pptx", "office", purpose="reading .pptx files")
        try:
            presentation = pptx.Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise LoadError(f"could not open presentation {path}: {exc}") from exc

        blocks: list[TextBlock] = []
        notes_count = 0

        for slide_number, slide in enumerate(presentation.slides, start=1):
            locator = f"slide {slide_number}"
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = " ".join(shape.text_frame.text.split())
                    if text:
                        blocks.append(
                            TextBlock(kind=BLOCK_SLIDE, text=text, locator=locator)
                        )
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [" ".join(c.text.split()) for c in row.cells]
                        blocks.append(
                            TextBlock(
                                kind=BLOCK_TABLE_ROW,
                                text=" | ".join(c for c in cells if c),
                                locator=locator,
                            )
                        )
            if slide.has_notes_slide:
                notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
                if notes:
                    notes_count += 1
                    blocks.append(
                        TextBlock(
                            kind=BLOCK_SLIDE, text=notes, locator=f"{locator} notes"
                        )
                    )

        return Document(
            source=source,
            text="\n\n".join(b.text for b in blocks),
            blocks=blocks,
            metadata={
                "slide_count": len(presentation.slides._sldIdLst),  # noqa: SLF001
                "notes_count": notes_count,
                **_pptx_properties(presentation),
            },
        )


def _pptx_properties(presentation: Any) -> dict[str, Any]:
    props = getattr(presentation, "core_properties", None)
    if props is None:
        return {}
    result: dict[str, Any] = {}
    for key in ("title", "author", "subject", "category"):
        value = getattr(props, key, None)
        if value:
            result[f"pptx_{key}"] = str(value)
    return result


class XlsxLoader:
    """Loader for Excel workbooks.

    Only the raw cell grid is recovered. Formula evaluation and type
    interpretation are intentionally out of scope here — that belongs in an
    extractor, where it can be validated against a contract.
    """

    name = "xlsx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    def __init__(self, *, max_rows_per_sheet: int = 20_000) -> None:
        self.max_rows_per_sheet = max_rows_per_sheet

    def load(self, path: Path, source: SourceRef) -> Document:
        openpyxl = require("openpyxl", "office", purpose="reading .xlsx files")
        try:
            workbook = openpyxl.load_workbook(
                str(path), read_only=True, data_only=True
            )
        except Exception as exc:  # noqa: BLE001
            raise LoadError(f"could not open workbook {path}: {exc}") from exc

        blocks: list[TextBlock] = []
        sheet_summaries: list[dict[str, Any]] = []

        try:
            for sheet in workbook.worksheets:
                row_count = 0
                header: list[str] = []
                for row_number, row in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    if row_number > self.max_rows_per_sheet:
                        break
                    values = ["" if v is None else str(v) for v in row]
                    if not any(v.strip() for v in values):
                        continue
                    if not header:
                        header = [v for v in values if v.strip()]
                    row_count += 1
                    blocks.append(
                        TextBlock(
                            kind=BLOCK_SHEET_ROW,
                            text=" | ".join(v for v in values if v.strip()),
                            locator=f"{sheet.title}!{row_number}",
                        )
                    )
                sheet_summaries.append(
                    {
                        "title": sheet.title,
                        "rows": row_count,
                        "columns": header,
                    }
                )
        finally:
            workbook.close()

        return Document(
            source=source,
            text="\n".join(b.text for b in blocks),
            blocks=blocks,
            metadata={
                "sheet_count": len(sheet_summaries),
                "sheets": sheet_summaries,
                "row_count": len(blocks),
            },
        )
